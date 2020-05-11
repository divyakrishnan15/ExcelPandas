from openpyxl import load_workbook
from openpyxl.chart import BarChart,PieChart,Series,Reference

wb=load_workbook('crime_report_output.xlsx')
ws=wb.active

chart=BarChart()
data=Reference(ws,min_row=8,min_col=1,max_col=13,max_row=13)
chart.add_data(data,titles_from_data=True)

ws.add_chart(chart,'B14')
wb.save('lines.xlsx')

